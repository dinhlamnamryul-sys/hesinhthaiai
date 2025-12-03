import os
import io
import math
import json
import tempfile
import streamlit as st
import requests
from typing import Tuple, List

# --- Xử lý Import Error ---
# Phần này giúp báo lỗi rõ ràng trên giao diện nếu thiếu thư viện
try:
    import pdfplumber
    from docx import Document
    from bs4 import BeautifulSoup
except ImportError as e:
    st.error(f"Lỗi thiếu thư viện: {e}")
    st.info("Vui lòng đảm bảo file 'requirements.txt' đã có đầy đủ: pdfplumber, requests, python-docx, beautifulsoup4")
    st.stop()

# ------------------------- CONFIG -------------------------
st.set_page_config(page_title="Tạo đề & Ma trận (Gemini AI)", page_icon="📝", layout="wide")
st.title("📝 Tạo ma trận & đề kiểm tra (Gemini AI) — upload sách, công văn → AI trả về ma trận & đề")

# Lấy API Key từ Secrets hoặc biến môi trường
if "GOOGLE_API_KEY" in st.secrets:
    GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
else:
    GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")

# Nếu chưa có Key, cho nhập tay
if not GOOGLE_API_KEY:
    with st.expander("⚠️ Chưa cấu hình API Key", expanded=True):
        GOOGLE_API_KEY = st.text_input("Nhập Google API Key của bạn:", type="password")
        st.markdown("[Lấy API Key miễn phí tại đây](https://aistudio.google.com/app/apikey)")

if not GOOGLE_API_KEY:
    st.stop()

# ------------------------- HELPERS: TẬP TIN -> TEXT -------------------------
def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception as e:
        st.warning(f"Không thể đọc PDF bình thường: {e}.")
    return "\n".join(text_parts)

def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        return "\n".join(paragraphs)
    except Exception:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(file_bytes)
            tmp.flush()
            try:
                doc = Document(tmp.name)
                paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
                return "\n".join(paragraphs)
            finally:
                os.unlink(tmp.name)

def extract_text_from_file(uploaded) -> Tuple[str, str]:
    if uploaded is None:
        return ("", "")
    raw = uploaded.read()
    uploaded.seek(0)
    name_lower = uploaded.name.lower()
    
    if name_lower.endswith(".pdf"):
        return ("application/pdf", extract_text_from_pdf(raw))
    if name_lower.endswith(".docx"):
        return ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", extract_text_from_docx(raw))
    if name_lower.endswith(".doc"):
        try:
            return ("application/msword", extract_text_from_docx(raw))
        except Exception:
            return ("application/msword", raw.decode(errors="ignore"))
    try:
        return ("text/plain", raw.decode("utf-8"))
    except Exception:
        return ("application/octet-stream", raw.decode(errors="ignore"))

# ------------------------- HELPERS: GỌI GEMINI API -------------------------
def call_gemini_generate_matrix_and_exam(api_key: str, textbook_text: str, official_doc_text: str, template_text: str, instruction: str) -> dict:
    """
    Gọi Google Gemini API để sinh ma trận và đề thi dưới dạng JSON.
    Sử dụng model gemini-2.5-flash cho tốc độ nhanh và context lớn.
    """
    # Gemini Flash có context window rất lớn (1M token), nên ta có thể gửi nhiều text hơn mà không cần cắt quá nhỏ.
    # Tuy nhiên, vẫn nên giới hạn để tránh timeout hoặc lỗi quá tải nếu file quá khổng lồ.
    MAX_CHARS = 200000 # Khoảng 50k token, dư sức cho hầu hết SGK chương/bài
    
    if len(textbook_text) > MAX_CHARS:
        textbook_text = textbook_text[:MAX_CHARS] + "\n...(đã cắt bớt)..."
    
    system_msg = (
        "Bạn là một chuyên gia giáo dục chuyên tạo MA TRẬN (dạng bảng HTML) và ĐỀ KIỂM TRA (HTML) "
        "theo đúng MẪU đề được cung cấp. "
        "Nhiệm vụ của bạn là trả về kết quả dưới dạng JSON hợp lệ."
    )

    user_msg = (
        "Dưới đây là tài liệu nguồn:\n\n"
        f"=== NỘI DUNG SGK (Kiến thức nguồn) ===\n{textbook_text}\n\n"
        f"=== CÔNG VĂN / KHUNG CHƯƠNG TRÌNH ===\n{official_doc_text}\n\n"
        f"=== MẪU ĐỀ (Template Format) ===\n{template_text}\n\n"
        f"=== YÊU CẦU CỦA GIÁO VIÊN ===\n{instruction}\n\n"
        "Hãy thực hiện:\n"
        "1. Xây dựng MA TRẬN đề thi (matrixHtml) phù hợp với công văn và yêu cầu.\n"
        "2. Soạn ĐỀ THI (examHtml) dựa trên ma trận vừa tạo. Nội dung câu hỏi lấy từ SGK. Hình thức trình bày giống Mẫu Đề.\n"
        "Output JSON schema: { \"matrixHtml\": \"string (html code)\", \"examHtml\": \"string (html code)\" }"
    )

    # Cấu hình gọi API Gemini
    model = "gemini-2.5-flash" 
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
    
    headers = {
        "Content-Type": "application/json"
    }
    
    data = {
        "contents": [{
            "parts": [{"text": user_msg}]
        }],
        "systemInstruction": {
            "parts": [{"text": system_msg}]
        },
        "generationConfig": {
            "responseMimeType": "application/json", # Bắt buộc trả về JSON
            "temperature": 0.3
        }
    }

    try:
        response = requests.post(url, headers=headers, json=data, timeout=120)
        
        if response.status_code != 200:
            raise RuntimeError(f"Lỗi API ({response.status_code}): {response.text}")
            
        result_json = response.json()
        
        # Parse kết quả
        try:
            candidates = result_json.get("candidates", [])
            if not candidates:
                 raise RuntimeError("AI không trả về kết quả (No candidates).")
            
            content_text = candidates[0].get("content", {}).get("parts", [])[0].get("text", "")
            parsed = json.loads(content_text)
            return parsed
            
        except (KeyError, IndexError, json.JSONDecodeError) as e:
            raise RuntimeError(f"Lỗi xử lý dữ liệu trả về từ AI: {e}\nRaw: {result_json}")

    except requests.exceptions.Timeout:
        raise RuntimeError("Yêu cầu hết thời gian chờ (Timeout). Vui lòng thử lại.")
    except Exception as e:
        raise RuntimeError(f"Lỗi kết nối: {e}")

# ------------------------- HELPERS: HTML -> DOCX -------------------------
def html_to_plain_text(html: str) -> str:
    if not html:
        return ""
    soup = BeautifulSoup(html, "html.parser")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    return "\n".join(lines)

def make_docx_from_htmls(matrix_html: str, exam_html: str) -> bytes:
    doc = Document()
    doc.add_heading("KẾT QUẢ TẠO ĐỀ TỰ ĐỘNG", level=0)
    
    doc.add_heading("I. MA TRẬN ĐỀ THI", level=1)
    if matrix_html:
        matrix_text = html_to_plain_text(matrix_html)
        doc.add_paragraph(matrix_text)
    else:
        doc.add_paragraph("[Không có nội dung ma trận]")

    doc.add_page_break()
    
    doc.add_heading("II. ĐỀ KIỂM TRA", level=1)
    if exam_html:
        exam_text = html_to_plain_text(exam_html)
        doc.add_paragraph(exam_text)
    else:
        doc.add_paragraph("[Không có nội dung đề]")
        
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.read()

# ------------------------- STREAMLIT UI -------------------------
st.info("💡 Ứng dụng sử dụng **Google Gemini 2.5 Flash**. Vui lòng nhập API Key để bắt đầu.")

col1, col2 = st.columns(2)
with col1:
    uploaded_textbook = st.file_uploader("1. Sách giáo khoa (Nguồn kiến thức)", type=['pdf','docx','doc'], key='tb')
    uploaded_official = st.file_uploader("2. Công văn / Khung chương trình", type=['pdf','docx','doc'], key='cv')

with col2:
    uploaded_template = st.file_uploader("3. Mẫu đề kiểm tra (Format)", type=['pdf','docx','doc'], key='tpl')
    instruction = st.text_area("Yêu cầu cụ thể (Số câu, tỉ lệ, mức độ...)", 
                               value="Tạo ma trận 21 câu (Trắc nghiệm 5đ, Tự luận 5đ). Tỉ lệ NB/TH/VD: 40/30/30.", 
                               height=120)

if st.button("🚀 TẠO MA TRẬN & ĐỀ", type="primary"):
    if not GOOGLE_API_KEY:
         st.error("⚠️ Vui lòng nhập Google API Key.")
    elif not uploaded_textbook or not uploaded_official or not uploaded_template:
        st.error("⚠️ Vui lòng tải lên đủ 3 file: SGK, Công văn, Mẫu đề.")
    else:
        with st.status("Đang xử lý với Gemini AI...", expanded=True) as status:
            st.write("📖 Đang đọc nội dung file...")
            tb_mime, tb_text = extract_text_from_file(uploaded_textbook)
            cv_mime, cv_text = extract_text_from_file(uploaded_official)
            tpl_mime, tpl_text = extract_text_from_file(uploaded_template)
            
            st.write(f"✅ Đã đọc xong: SGK ({len(tb_text)} ký tự), Công văn ({len(cv_text)} ký tự).")
            
            st.write("🤖 Đang gửi dữ liệu cho Gemini phân tích...")
            try:
                result = call_gemini_generate_matrix_and_exam(GOOGLE_API_KEY, tb_text, cv_text, tpl_text, instruction)
                status.update(label="Hoàn tất!", state="complete", expanded=False)
            except Exception as e:
                status.update(label="Gặp lỗi!", state="error")
                st.error(f"Lỗi trong quá trình xử lý: {e}")
                st.stop()

        # Validate result
        matrix_html = result.get("matrixHtml") or result.get("matrix") or ""
        exam_html = result.get("examHtml") or result.get("exam") or ""

        if not matrix_html and not exam_html:
            st.error("AI không trả về kết quả đúng định dạng.")
            with st.expander("Xem dữ liệu trả về từ AI"):
                st.write(result)
            st.stop()

        # Hiển thị kết quả
        tab1, tab2 = st.tabs(["📊 Ma trận", "📝 Đề kiểm tra"])
        
        with tab1:
            st.markdown(matrix_html, unsafe_allow_html=True)
            st.download_button("📥 Tải HTML Ma trận", matrix_html, "matran.html", "text/html")
            
        with tab2:
            st.markdown(exam_html, unsafe_allow_html=True)
            st.download_button("📥 Tải HTML Đề", exam_html, "de_kiem_tra.html", "text/html")

        # Tải DOCX chung
        docx_bytes = make_docx_from_htmls(matrix_html, exam_html)
        st.markdown("---")
        st.download_button(
            label="📥 TẢI VỀ FILE WORD (.DOCX)",
            data=docx_bytes,
            file_name="De_Kiem_Tra_Gemini_Generated.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary"
        )
